import os
from collections import defaultdict

from openpyxl import Workbook, load_workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

DIRECTORIO = os.path.dirname(__file__)
ARCHIVO_EXCEL = os.path.join(DIRECTORIO, "empleados.xlsx")
ARCHIVO_PDF = os.path.join(DIRECTORIO, "reporte_empleados.pdf")

DATOS_EJEMPLO = [
    ("Ana Torres", "Ventas", 3200),
    ("Luis Ramirez", "Ventas", 2900),
    ("Maria Lopez", "Tecnologia", 4500),
    ("Carlos Diaz", "Tecnologia", 4800),
    ("Sofia Vargas", "Tecnologia", 5100),
    ("Jorge Castro", "Recursos Humanos", 3000),
    ("Elena Rios", "Recursos Humanos", 3100),
    ("Pedro Salas", "Finanzas", 3900),
    ("Lucia Mendez", "Finanzas", 4100),
    ("Diego Herrera", "Ventas", 3050),
]


def crear_archivo_ejemplo(ruta):
    libro = Workbook()
    hoja = libro.active
    hoja.append(["nombre", "departamento", "salario"])
    for fila in DATOS_EJEMPLO:
        hoja.append(fila)
    libro.save(ruta)


def leer_empleados(ruta):
    libro = load_workbook(ruta)
    hoja = libro.active
    empleados = []
    for nombre, departamento, salario in hoja.iter_rows(min_row=2, values_only=True):
        if nombre is None:
            continue
        empleados.append({"nombre": nombre, "departamento": departamento, "salario": float(salario)})
    return empleados


def calcular_estadisticas_por_departamento(empleados):
    salarios_por_departamento = defaultdict(list)
    for empleado in empleados:
        salarios_por_departamento[empleado["departamento"]].append(empleado["salario"])

    estadisticas = {}
    for departamento, salarios in salarios_por_departamento.items():
        estadisticas[departamento] = {
            "total_empleados": len(salarios),
            "salario_promedio": sum(salarios) / len(salarios),
        }
    return estadisticas


def generar_pdf(estadisticas, total_empleados, salario_promedio_general, ruta_salida):
    documento = SimpleDocTemplate(ruta_salida, pagesize=letter)
    estilos = getSampleStyleSheet()
    elementos = []

    elementos.append(Paragraph("Reporte de Empleados por Departamento", estilos["Title"]))
    elementos.append(Spacer(1, 0.5 * cm))

    resumen = (
        f"Este reporte resume la plantilla de la empresa, compuesta por {total_empleados} "
        f"empleados distribuidos en {len(estadisticas)} departamentos, con un salario "
        f"promedio general de ${salario_promedio_general:,.2f}. A continuacion se detalla "
        f"el total de empleados y el salario promedio por cada departamento."
    )
    elementos.append(Paragraph("Resumen Ejecutivo", estilos["Heading2"]))
    elementos.append(Paragraph(resumen, estilos["BodyText"]))
    elementos.append(Spacer(1, 0.7 * cm))

    elementos.append(Paragraph("Detalle por Departamento", estilos["Heading2"]))

    datos_tabla = [["Departamento", "Total de Empleados", "Salario Promedio"]]
    for departamento in sorted(estadisticas):
        stats = estadisticas[departamento]
        datos_tabla.append(
            [departamento, str(stats["total_empleados"]), f"${stats['salario_promedio']:,.2f}"]
        )

    tabla = Table(datos_tabla, colWidths=[7 * cm, 5 * cm, 5 * cm])
    tabla.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2c3e50")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("ALIGN", (1, 0), (-1, -1), "CENTER"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f2f2f2")]),
            ]
        )
    )
    elementos.append(tabla)

    documento.build(elementos)


if __name__ == "__main__":
    if not os.path.exists(ARCHIVO_EXCEL):
        crear_archivo_ejemplo(ARCHIVO_EXCEL)
        print(f"Archivo {ARCHIVO_EXCEL} no existía, se creó con datos de ejemplo.")

    empleados = leer_empleados(ARCHIVO_EXCEL)
    estadisticas = calcular_estadisticas_por_departamento(empleados)

    total_empleados = len(empleados)
    salario_promedio_general = sum(e["salario"] for e in empleados) / total_empleados

    generar_pdf(estadisticas, total_empleados, salario_promedio_general, ARCHIVO_PDF)
    print(f"Reporte generado en: {ARCHIVO_PDF}")
