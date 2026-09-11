import os
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from openpyxl import Workbook, load_workbook

DIRECTORIO = os.path.dirname(__file__)
ARCHIVO_EXCEL = os.path.join(DIRECTORIO, "reuniones.xlsx")
DIRECTORIO_SALIDA = os.path.join(DIRECTORIO, "agendas")

DATOS_EJEMPLO = [
    ("Kickoff Proyecto Alfa", "Ana Torres, Luis Ramirez, Maria Lopez", 60, "Alcance y cronograma del proyecto"),
    ("Revision Sprint 4", "Carlos Diaz, Sofia Vargas", 45, "Resultados del sprint y bloqueos"),
    ("Comite de Presupuesto", "Jorge Castro, Elena Rios, Pedro Salas, Lucia Mendez", 90, "Aprobacion del presupuesto anual"),
]

# Proporciones estandar de una agenda: apertura, tema clave, discusion, acuerdos, cierre
PROPORCIONES_AGENDA = [
    ("Apertura y bienvenida", 0.10),
    ("Presentacion del tema clave", 0.45),
    ("Discusion y aportes de los participantes", 0.25),
    ("Acuerdos y proximos pasos", 0.15),
    ("Cierre", 0.05),
]


def crear_archivo_ejemplo(ruta):
    libro = Workbook()
    hoja = libro.active
    hoja.append(["nombre_reunion", "participantes", "duracion_minutos", "tema_clave"])
    for fila in DATOS_EJEMPLO:
        hoja.append(fila)
    libro.save(ruta)


def leer_reuniones(ruta):
    libro = load_workbook(ruta)
    hoja = libro.active
    reuniones = []
    for nombre, participantes, duracion, tema_clave in hoja.iter_rows(min_row=2, values_only=True):
        if nombre is None:
            continue
        reuniones.append(
            {
                "nombre": nombre,
                "participantes": [p.strip() for p in str(participantes).split(",")],
                "duracion_minutos": int(duracion),
                "tema_clave": tema_clave,
            }
        )
    return reuniones


def generar_orden_del_dia(tema_clave, duracion_minutos):
    orden_del_dia = []
    minutos_asignados = 0
    for indice, (punto, proporcion) in enumerate(PROPORCIONES_AGENDA):
        if indice == len(PROPORCIONES_AGENDA) - 1:
            minutos = duracion_minutos - minutos_asignados
        else:
            minutos = round(duracion_minutos * proporcion)
            minutos_asignados += minutos

        if punto == "Presentacion del tema clave":
            punto = f"Presentacion del tema clave: {tema_clave}"

        orden_del_dia.append((punto, minutos))

    return orden_del_dia


def nombre_archivo_seguro(nombre_reunion):
    nombre = re.sub(r"[^\w\s-]", "", nombre_reunion).strip().replace(" ", "_")
    return f"{nombre}.docx"


def generar_documento_reunion(reunion, ruta_salida):
    documento = Document()

    # Portada
    titulo = documento.add_heading(reunion["nombre"], level=0)
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitulo = documento.add_paragraph("Agenda de Reunion")
    subtitulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitulo.runs[0].font.size = Pt(16)

    documento.add_paragraph()
    documento.add_paragraph(f"Duracion estimada: {reunion['duracion_minutos']} minutos")
    documento.add_paragraph(f"Tema clave: {reunion['tema_clave']}")
    documento.add_paragraph("Participantes: " + ", ".join(reunion["participantes"]))

    documento.add_page_break()

    # Orden del dia
    documento.add_heading("Orden del Dia", level=1)
    orden_del_dia = generar_orden_del_dia(reunion["tema_clave"], reunion["duracion_minutos"])

    tabla = documento.add_table(rows=1, cols=2)
    tabla.style = "Light Grid Accent 1"
    encabezados = tabla.rows[0].cells
    encabezados[0].text = "Punto de la agenda"
    encabezados[1].text = "Tiempo asignado"

    for punto, minutos in orden_del_dia:
        fila = tabla.add_row().cells
        fila[0].text = punto
        fila[1].text = f"{minutos} min"

    documento.add_paragraph()

    # Observaciones
    documento.add_heading("Observaciones", level=1)
    documento.add_paragraph(
        "Espacio para registrar comentarios, decisiones adicionales o pendientes surgidos "
        "durante la reunion."
    )
    for _ in range(3):
        documento.add_paragraph("_" * 80)

    documento.save(ruta_salida)


if __name__ == "__main__":
    if not os.path.exists(ARCHIVO_EXCEL):
        crear_archivo_ejemplo(ARCHIVO_EXCEL)
        print(f"Archivo {ARCHIVO_EXCEL} no existía, se creó con datos de ejemplo.")

    os.makedirs(DIRECTORIO_SALIDA, exist_ok=True)

    reuniones = leer_reuniones(ARCHIVO_EXCEL)
    for reunion in reuniones:
        ruta_salida = os.path.join(DIRECTORIO_SALIDA, nombre_archivo_seguro(reunion["nombre"]))
        generar_documento_reunion(reunion, ruta_salida)
        print(f"Agenda generada: {ruta_salida}")
